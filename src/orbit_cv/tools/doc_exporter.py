# src/orbit_cv/tools/doc_exporter.py
import re
from pathlib import Path
from typing import Literal
from docx import Document
from fpdf import FPDF
from fpdf.errors import FPDFException
from langchain_core.tools import tool

from orbit_cv.paths import EXPORTS_DIR, resolve_path


def _sanitize_utf8_for_pdf(text: str) -> str:
    """Replaces Unicode characters outside latin-1 (em-dashes, smart quotes, bullet symbols)
    with ASCII equivalents compatible with standard FPDF Helvetica.
    """
    replacements = {
        "—": " - ",  # Em-dash
        "–": "-",    # En-dash
        "“": '"',    # Left double quote
        "”": '"',    # Right double quote
        "‘": "'",    # Left single quote
        "’": "'",    # Right single quote
        "•": "*",    # Bullet point
        "…": "...",  # Ellipsis
    }
    for orig, repl in replacements.items():
        text = text.replace(orig, repl)

    # Encode to latin-1 with replacement fallback for any remaining unsupported chars
    return text.encode("latin-1", "replace").decode("latin-1")


def _sanitize_token_lengths(text: str, max_len: int = 40) -> str:
    """Inserts soft spaces into ultra-long unbroken words/URLs so FPDF line breaks won't crash."""
    words = text.split(" ")
    processed = []
    for word in words:
        if len(word) > max_len:
            chunks = [word[i : i + max_len] for i in range(0, len(word), max_len)]
            processed.append(" ".join(chunks))
        else:
            processed.append(word)
    return " ".join(processed)


def _sanitize_filename(prefix: str) -> str:
    safe = re.sub(r"[^\w\-_]", "", prefix.strip().replace(" ", "_"))
    return safe if safe else "tailored_document"


class CleanPDF(FPDF):

    def header(self):
        self.set_font("Helvetica", "B", 8)
        self.set_text_color(120, 120, 120)
        self.cell(0, 5, _sanitize_utf8_for_pdf("Orbit - Career Accelerator Document"), align="R")
        self.ln(6)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "I", 8)
        self.set_text_color(150, 150, 150)
        self.cell(0, 10, f"Page {self.page_no()}", align="C")


@tool
def export_document(
    content_markdown: str,
    output_format: Literal["pdf", "docx", "md"] = "pdf",
    filename_prefix: str = "tailored_cv",
) -> str:
    """Exports tailored resumes, cover letters, or career strategy plans into styled PDF, DOCX, or MD files.

    Returns the absolute path to the generated document file.
    """
    print(f"tool call [export_document] - content_markdown length: {len(content_markdown)}")
    safe_prefix = _sanitize_filename(filename_prefix)

    # Ensure output target directory exists in data/exports
    EXPORTS_DIR.mkdir(parents=True, exist_ok=True)

    # 1. Export Markdown
    if output_format == "md":
        out_path = resolve_path(f"{safe_prefix}.md", default_route="exports")
        out_path.write_text(content_markdown, encoding="utf-8")
        return str(out_path.resolve())

    # 2. Export DOCX
    elif output_format == "docx":
        out_path = resolve_path(f"{safe_prefix}.docx", default_route="exports")
        doc = Document()

        for line in content_markdown.split("\n"):
            line_str = line.strip()
            if not line_str:
                continue

            if line_str.startswith("# "):
                doc.add_heading(line_str[2:], level=1)
            elif line_str.startswith("## "):
                doc.add_heading(line_str[3:], level=2)
            elif line_str.startswith("### "):
                doc.add_heading(line_str[4:], level=3)
            elif line_str.startswith("- ") or line_str.startswith("* "):
                doc.add_paragraph(line_str[2:], style="List Bullet")
            elif line_str.startswith("---"):
                doc.add_paragraph("____________________________________________________")
            else:
                doc.add_paragraph(line_str)

        doc.save(out_path)
        return str(out_path.resolve())

    # 3. Export PDF
    elif output_format == "pdf":
        out_path = resolve_path(f"{safe_prefix}.pdf", default_route="exports")
        pdf = CleanPDF(format="A4", unit="mm")

        # Explicit margins ensure effective printable width is always bounded (~190mm)
        pdf.set_margins(left=10, top=10, right=10)
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()

        def _safe_multi_cell(text: str, h: int = 5):
            effective_width = pdf.epw
            try:
                pdf.multi_cell(w=effective_width, h=h, text=text)
            except FPDFException:
                # Fallback truncate in case of unparseable layout state
                pdf.multi_cell(w=effective_width, h=h, text=text[:100] + "...")

        for line in content_markdown.split("\n"):
            line_s = line.strip()
            if not line_s:
                pdf.ln(3)
                continue

            # Double sanitize: convert unicode to Latin-1 and break ultra-long tokens/URLs
            clean_line = _sanitize_token_lengths(_sanitize_utf8_for_pdf(line_s))

            # Title / H1
            if clean_line.startswith("# "):
                pdf.set_font("Helvetica", "B", 16)
                pdf.set_text_color(15, 23, 42)
                _safe_multi_cell(clean_line[2:], h=8)
                pdf.ln(2)

            # Subtitle / H2
            elif clean_line.startswith("## "):
                pdf.set_font("Helvetica", "B", 12)
                pdf.set_text_color(37, 99, 235)
                _safe_multi_cell(clean_line[3:].upper(), h=7)
                pdf.ln(1)

            # Section Header / H3
            elif clean_line.startswith("### "):
                pdf.set_font("Helvetica", "B", 10)
                pdf.set_text_color(30, 41, 59)
                _safe_multi_cell(clean_line[4:], h=6)

            # Bullet points
            elif clean_line.startswith("- ") or clean_line.startswith("* "):
                pdf.set_font("Helvetica", size=9.5)
                pdf.set_text_color(51, 65, 85)
                bullet_text = clean_line[2:]
                _safe_multi_cell(f"  *  {bullet_text}", h=5)

            # Horizontal separator
            elif clean_line.startswith("---"):
                pdf.set_draw_color(226, 232, 240)
                pdf.line(pdf.get_x(), pdf.get_y(), pdf.get_x() + pdf.epw, pdf.get_y())
                pdf.ln(4)

            # Regular body paragraph
            else:
                pdf.set_font("Helvetica", size=9.5)
                pdf.set_text_color(51, 65, 85)
                _safe_multi_cell(clean_line, h=5)

        pdf.output(str(out_path))
        return str(out_path.resolve())

    else:
        raise ValueError(f"Unsupported output format: {output_format}")